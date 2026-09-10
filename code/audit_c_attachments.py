"""Read-only audit of C-problem inputs; writes derived JSON/CSV, never XLSX."""
from pathlib import Path
from datetime import datetime, timedelta, time
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'CUMCM2026Problems' / 'C题' / '附件'
OUT = ROOT / 'results' / 'attachment_audit'


def clock_minutes(v):
    if isinstance(v, (time, datetime)):
        return v.hour * 60 + v.minute
    text = str(v)
    offset = 1440 if '+1' in text else 0
    h, m = text.replace('+1', '').split(':')[:2]
    return int(h) * 60 + int(m) + offset


def matrix(path, sheet=0):
    frame = pd.read_excel(path, sheet_name=sheet)
    dates = pd.to_datetime(frame.iloc[:, 0])
    times = [clock_minutes(v) for v in frame.columns[1:]]
    values = frame.iloc[:, 1:].apply(pd.to_numeric, errors='raise').to_numpy(float)
    assert times == list(range(10, 1441, 10)), times
    assert dates.tolist() == pd.date_range('2025-01-01', '2025-12-31').tolist()
    assert values.shape == (365, 144)
    return dates, values


def stats(a):
    return dict(count=int(a.size), missing=int(np.isnan(a).sum()),
                nonfinite=int((~np.isfinite(a)).sum()), minimum=float(np.nanmin(a)),
                maximum=float(np.nanmax(a)), mean=float(np.nanmean(a)),
                negative=int((a < 0).sum()), zero=int((a == 0).sum()))


def scores(y, pred):
    error = np.asarray(pred) - np.asarray(y)
    return dict(n=int(error.size), mae_kw=float(np.abs(error).mean()),
                rmse_kw=float(np.sqrt(np.square(error).mean())),
                bias_kw=float(error.mean()))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    inventory = []
    for path in sorted(SOURCE.rglob('*.xlsx')):
        wb = load_workbook(path, data_only=False)
        book = dict(file=str(path.relative_to(SOURCE)), sha256=hashlib.sha256(path.read_bytes()).hexdigest(), sheets=[])
        for sheet in wb:
            comments, formulas, errors = [], [], []
            for row in sheet:
                for cell in row:
                    if cell.comment:
                        comments.append(dict(cell=cell.coordinate, text=cell.comment.text))
                    if cell.data_type == 'f':
                        formulas.append(dict(cell=cell.coordinate, formula=cell.value))
                    if cell.data_type == 'e':
                        errors.append(dict(cell=cell.coordinate, value=cell.value))
            book['sheets'].append(dict(name=sheet.title, rows=sheet.max_row, columns=sheet.max_column,
                headers=[str(c.value) if c.value is not None else None for c in sheet[1]],
                merged=[str(r) for r in sheet.merged_cells.ranges], comments=comments,
                formula_count=len(formulas), formula_examples=formulas[:10], errors=errors,
                nonempty_first_column=[dict(row=i, value=str(sheet.cell(i,1).value)) for i in range(1,sheet.max_row+1) if sheet.cell(i,1).value is not None][:40]))
        inventory.append(book)
        wb.close()

    one = pd.read_excel(SOURCE / '附件1.xlsx')
    assert [clock_minutes(x) for x in one.iloc[:,0]] == list(range(10,1441,10))
    p1, l1, v1 = (one.iloc[:,i].to_numpy(float) for i in [1,2,3])
    dates, load = matrix(SOURCE / '附件2.xlsx', '小区负载')
    _, pv = matrix(SOURCE / '附件2.xlsx', '光伏发电实际功率')
    _, price = matrix(SOURCE / '附件4.xlsx')
    # Energy figures use a declared interval-average interpretation of the supplied powers.
    daily = pd.DataFrame(dict(date=dates, load_kwh=load.sum(axis=1)/6,
        pv_kwh=pv.sum(axis=1)/6, net_import_no_storage_kwh=np.maximum(load-pv,0).sum(axis=1)/6,
        surplus_no_storage_kwh=np.maximum(pv-load,0).sum(axis=1)/6,
        fixed_tariff_no_storage_yuan=(np.maximum(load-pv,0)*p1).sum(axis=1)/6,
        varying_tariff_no_storage_yuan=(np.maximum(load-pv,0)*price).sum(axis=1)/6,
        price_mean=price.mean(axis=1), price_min=price.min(axis=1), price_max=price.max(axis=1)))
    daily.to_csv(OUT / 'daily_descriptive_statistics.csv', index=False, encoding='utf-8-sig')
    monthly = daily.groupby(daily.date.dt.month).agg({'load_kwh':'mean','pv_kwh':'mean',
        'surplus_no_storage_kwh':'mean','price_mean':'mean','price_min':'min','price_max':'max'})
    monthly.to_csv(OUT / 'monthly_descriptive_statistics.csv', encoding='utf-8-sig')

    simple = []
    for label, a in [('load',load), ('pv',pv), ('net_load',load-pv), ('price',price)]:
        actual = a[31:]
        preds = {
            'previous_day': a[30:-1],
            'previous_week': a[24:-7],
            'past_28day_slot_mean': np.array([a[max(0,i-28):i].mean(axis=0) for i in range(31,365)]),
            'past_4_same_weekday_mean': np.array([a[[i-j for j in [7,14,21,28] if i-j>=0]].mean(axis=0) for i in range(31,365)]),
        }
        if label == 'price':
            preds['attachment1_fixed_curve'] = np.broadcast_to(p1, actual.shape)
        for method, prediction in preds.items():
            result = dict(series=label, method=method, **scores(actual,prediction))
            if label == 'price':
                result = {k.replace('_kw','_yuan_per_kwh'):v for k,v in result.items()}
            if label == 'pv':
                mask=actual>50
                result['daylight_actual_gt50'] = scores(actual[mask],prediction[mask])
            simple.append(result)

    raw = pd.read_excel(SOURCE / '附件3.xlsx')
    issue_days = pd.to_datetime(raw.iloc[:,0].replace('',np.nan).ffill())
    issue_hours = [clock_minutes(v)//60 for v in raw.iloc[:,1]]
    forecast_values = raw.iloc[:,2:].apply(pd.to_numeric,errors='raise').to_numpy(float)
    assert forecast_values.shape == (1460,24)
    assert issue_hours == [0,6,12,18]*365
    assert issue_days.tolist() == list(pd.date_range('2025-01-01','2025-12-31').repeat(4))
    lookup = {}
    records = []
    start = datetime(2025,1,1)
    for i,(day,hour) in enumerate(zip(issue_days,issue_hours)):
        issue = day.to_pydatetime()+timedelta(hours=hour)
        for lead in range(1,25):
            valid = issue+timedelta(hours=lead)
            value = float(forecast_values[i,lead-1])
            lookup[(issue,valid)] = value
            # Valid midnight belongs to the preceding day's final column.
            flat_index = int((valid-start).total_seconds()/600)-1
            actual = float(pv.ravel()[flat_index]) if 0 <= flat_index < pv.size else None
            records.append(dict(issue_time=issue,valid_time=valid,issue_hour=hour,
                lead_hours=lead, forecast_kw=value, actual_kw=actual))
    aligned=pd.DataFrame(records)
    aligned.to_csv(OUT / 'forecast_hourly_alignment.csv',index=False,encoding='utf-8-sig')
    evalf=aligned[(aligned.issue_time>=datetime(2025,2,1)) & aligned.actual_kw.notna()]
    forecasts=[]
    for hour,group in evalf.groupby('issue_hour'):
        mask=group.actual_kw>50
        forecasts.append(dict(issue_hour=int(hour),all_valid_hours=scores(group.actual_kw,group.forecast_kw),
            daylight_actual_gt50=scores(group.loc[mask,'actual_kw'],group.loc[mask,'forecast_kw'])))
    # Compare each updated forecast with 00:00 on exactly the same 6 subsequent hourly targets.
    comparison=[]
    for hour in [6,12,18]:
        yy,old,new=[],[],[]
        for day in pd.date_range('2025-02-01','2025-12-31'):
            issue=day.to_pydatetime()+timedelta(hours=hour)
            for step in range(1,7):
                valid=issue+timedelta(hours=step)
                flat_index=int((valid-start).total_seconds()/600)-1
                if not (0<=flat_index<pv.size):
                    continue
                yy.append(pv.ravel()[flat_index])
                old.append(lookup[(day.to_pydatetime(),valid)])
                new.append(lookup[(issue,valid)])
        y,o,n=map(np.asarray,[yy,old,new]); mask=y>50
        comparison.append(dict(update_hour=hour,original_00=scores(y,o),updated=scores(y,n),
            days_with_lower_mae=int((abs(n-y).reshape(-1,6).mean(axis=1)<abs(o-y).reshape(-1,6).mean(axis=1)-1e-9).sum()),
            daylight_original_00=scores(y[mask],o[mask]) if mask.any() else None,
            daylight_updated=scores(y[mask],n[mask]) if mask.any() else None))

    # At 18:00, count whether any supplied actual/forecast PV remains within that same calendar day.
    remaining=aligned[(aligned.issue_hour==18)&(aligned.lead_hours<=6)]
    report = dict(
        input_inventory=inventory,
        convention='Descriptive kWh values assume each supplied 10-minute power represents its interval average; point-to-interval interpretation remains an assumption.',
        attachment1=dict(price=stats(p1),load=stats(l1),pv=stats(v1),load_kwh=float(l1.sum()/6),
            pv_kwh=float(v1.sum()/6),surplus_no_storage_kwh=float(np.maximum(v1-l1,0).sum()/6),
            no_storage_purchase_kwh=float(np.maximum(l1-v1,0).sum()/6),
            no_storage_cost_yuan=float((np.maximum(l1-v1,0)*p1).sum()/6),
            tariff_hourly_mean=p1.reshape(24,6).mean(axis=1).tolist()),
        actuals=dict(load_kw=stats(load),pv_kw=stats(pv),price=stats(price),
            days_with_pv_surplus=int(((pv-load)>0).any(axis=1).sum()),
            surplus_interval_fraction=float((pv>load).mean()),
            duplicate_daily_load_profiles=int(pd.DataFrame(load).duplicated().sum()),
            duplicate_daily_pv_profiles=int(pd.DataFrame(pv).duplicated().sum()),
            price_correlation_fixed_curve=float(np.corrcoef(price.ravel(),np.tile(p1,365))[0,1]),
            formal_days=334,formal_intervals=334*144),
        baseline_forecast_metrics=simple,
        provided_forecast_stats=stats(forecast_values),
        forecast_unverifiable_targets=int(aligned.actual_kw.isna().sum()),
        provided_forecast_metrics=forecasts,
        matched_six_hour_update_comparison=comparison,
        remaining_today_after_18=dict(positive_forecast_count=int((remaining.forecast_kw>0).sum()),
            positive_actual_count=int((remaining.actual_kw>0).sum()),
            actual_1810_to_2400_energy_mean_kwh=float(pv[:,108:].sum(axis=1).mean()/6),
            actual_1810_to_2400_energy_max_kwh=float(pv[:,108:].sum(axis=1).max()/6)),
        attachment1_vs_full_year_mean={name:dict(max_abs=float(abs(a.mean(axis=0)-b).max()),
            mean_abs=float(abs(a.mean(axis=0)-b).mean())) for name,a,b in [('load_kw',load,l1),('pv_kw',pv,v1),('price_yuan_per_kwh',price,p1)]},
        mixed_net_forecast=scores((load-pv)[31:],load[24:-7]-pv[30:-1]),
        specified_days=daily.loc[daily.date.isin(pd.to_datetime(['2025-03-20','2025-06-21','2025-09-23','2025-12-21']))].to_dict('records'),
        monthly_descriptive=monthly.reset_index().to_dict('records'))
    (OUT / 'audit_summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    brief={k:v for k,v in report.items() if k not in ['input_inventory','monthly_descriptive']}
    print(json.dumps(brief,ensure_ascii=False,indent=2,default=str))


if __name__=='__main__':
    main()
